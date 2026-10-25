"""Markdown to DOCX conversion for the EmyDocx MVP."""

from io import BytesIO
import re

import mistune
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor


# XML 1.0 allows tab, line feed and carriage return, but not most C0
# control characters, isolated surrogates, or the FFFE/FFFF noncharacters.
ILLEGAL_XML_CHARS = re.compile(
    r"[^\x09\x0A\x0D\x20-\uD7FF\uE000-\uFFFD\U00010000-\U0010FFFF]"
)


def clean_xml_string(value) -> str:
    """Convert a value to text and remove characters forbidden in XML 1.0."""
    if value is None:
        return ""
    return ILLEGAL_XML_CHARS.sub("", str(value))


def _rgb(hex_color, fallback="#202124"):
    value = (hex_color or fallback).strip().lstrip("#")
    if len(value) != 6:
        value = fallback.lstrip("#")
    try:
        return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))
    except (ValueError, TypeError):
        return RGBColor(32, 33, 36)


def _apply_run_format(run, settings, heading=False, level=0):
    run.font.name = settings.get("font_name", "Arial")
    if heading:
        size_key = {1: "h1_size", 2: "h2_size", 3: "h3_size"}.get(level, "h3_size")
        run.font.size = Pt(float(settings.get(size_key, settings.get("body_size", 12))))
        run.font.color.rgb = _rgb(settings.get("heading_color", "#172554"), "#172554")
        run.bold = True
    else:
        run.font.size = Pt(float(settings.get("body_size", 12)))
        run.font.color.rgb = _rgb(settings.get("body_color", "#202124"), "#202124")


def _add_run(paragraph, text, settings, **formatting):
    """Add a run after sanitizing its text for XML compatibility."""
    run = paragraph.add_run(clean_xml_string(text))
    _apply_run_format(run, settings)
    for attribute, value in formatting.items():
        setattr(run, attribute, value)
    return run


def _add_inline(paragraph, text, settings, base_bold=False, base_italic=False):
    """Format common inline Markdown tokens in a text fragment."""
    text = clean_xml_string(text)
    pattern = re.compile(
        r"(\*\*\*.+?\*\*\*|___.+?___|\*\*.+?\*\*|__.+?__|"
        r"\*[^*\n]+?\*|_[^_\n]+?_|`[^`\n]+`|"
        r"\[([^\]]+)\]\(([^)]+)\))"
    )

    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            _add_run(
                paragraph, text[pos:match.start()], settings,
                bold=base_bold, italic=base_italic
            )

        token = match.group(0)
        if (token.startswith("***") and token.endswith("***")) or (
            token.startswith("___") and token.endswith("___")
        ):
            run = _add_run(paragraph, token[3:-3], settings, bold=True, italic=True)
        elif token.startswith("**") or token.startswith("__"):
            run = _add_run(paragraph, token[2:-2], settings, bold=True)
        elif token.startswith("*") or token.startswith("_"):
            run = _add_run(paragraph, token[1:-1], settings, italic=True)
        elif token.startswith("`"):
            run = _add_run(paragraph, token[1:-1], settings)
            run.font.name = "Consolas"
            run.font.color.rgb = _rgb("#7C3AED")
        else:
            label, url = match.group(2), match.group(3)
            run = _add_run(paragraph, f"{label} ({url})", settings)
            run.underline = True
        pos = match.end()

    if pos < len(text):
        _add_run(
            paragraph, text[pos:], settings,
            bold=base_bold, italic=base_italic
        )


def markdown_to_docx(markdown_text, settings):
    """Convert Markdown to a formatted DOCX byte string."""
    document = Document()
    section = document.sections[0]
    margin = Cm(float(settings.get("margin_cm", 2.0)))
    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin

    normal = document.styles["Normal"]
    normal.font.name = settings.get("font_name", "Arial")
    normal.font.size = Pt(float(settings.get("body_size", 12)))
    normal.font.color.rgb = _rgb(settings.get("body_color", "#202124"))
    normal.paragraph_format.line_spacing = float(settings.get("line_spacing", 1.15))
    normal.paragraph_format.space_after = Pt(6)

    # Enable AST support for Markdown tables, strikethrough, and LaTeX math.
    markdown = mistune.create_markdown(
        renderer="ast",
        plugins=["table", "strikethrough", "math"],
    )
    tokens = markdown(clean_xml_string(markdown_text or ""))

    def plain_text(children):
        parts = []
        for child in children or []:
            kind = child.get("type")
            if kind in ("text", "codespan", "inline_html", "inline_math", "block_math"):
                parts.append(child.get("raw", ""))
            elif kind in ("strong", "emphasis", "link", "strikethrough"):
                parts.append(plain_text(child.get("children")))
            elif kind == "image":
                parts.append(child.get("attrs", {}).get("alt", ""))
            elif kind == "softbreak":
                parts.append(" ")
            elif kind == "linebreak":
                parts.append("\n")
            elif kind == "table_cell":
                parts.append(plain_text(child.get("children", [])))
        return clean_xml_string("".join(parts))

    def inline_into(paragraph, children):
        for child in children or []:
            kind = child.get("type")
            raw = child.get("raw", "")

            if kind == "text":
                _add_inline(paragraph, raw, settings)
            elif kind == "strong":
                _add_run(paragraph, plain_text(child.get("children")), settings, bold=True)
            elif kind == "emphasis":
                _add_run(paragraph, plain_text(child.get("children")), settings, italic=True)
            elif kind == "codespan":
                run = _add_run(paragraph, raw, settings)
                run.font.name = "Consolas"
            elif kind == "link":
                label = plain_text(child.get("children"))
                url = child.get("attrs", {}).get("url", "")
                run = _add_run(paragraph, f"{label} ({url})" if url else label, settings)
                run.underline = True
            elif kind in ("softbreak", "linebreak"):
                paragraph.add_run("\n")
            elif kind == "image":
                alt = child.get("attrs", {}).get("alt", "Изображение")
                _add_run(paragraph, f"[{alt}]", settings)
            elif kind == "strikethrough":
                run = _add_run(paragraph, plain_text(child.get("children")), settings)
                run.font.strike = True
            elif kind == "inline_math":
                run = _add_run(paragraph, raw, settings)
                run.font.name = "Cambria Math"
                run.italic = True
            else:
                text = plain_text([child])
                if text:
                    _add_run(paragraph, text, settings)

    def add_table(block):
        """Render Mistune's table AST while keeping the header as one row."""
        parsed_rows = []
        header_rows = []
        for section_node in block.get("children", []):
            section_type = section_node.get("type")
            if section_type == "table_head":
                header_cells = section_node.get("children", [])
                header_rows.append([
                    plain_text(cell.get("children", [])) for cell in header_cells
                ])
            elif section_type == "table_body":
                for row_node in section_node.get("children", []):
                    if row_node.get("type") == "table_row":
                        parsed_rows.append([
                            plain_text(cell.get("children", []))
                            for cell in row_node.get("children", [])
                        ])
            elif section_type == "table_row":
                parsed_rows.append([
                    plain_text(cell.get("children", []))
                    for cell in section_node.get("children", [])
                ])

        all_rows = header_rows + parsed_rows
        if not all_rows:
            return

        column_count = max((len(row) for row in all_rows), default=0)
        if column_count == 0:
            return

        table = document.add_table(rows=0, cols=column_count)
        table.style = "Table Grid"

        for row_index, values in enumerate(all_rows):
            values = values + [""] * (column_count - len(values))
            cells = table.add_row().cells
            for col_index, value in enumerate(values):
                cell = cells[col_index]
                cell.text = ""
                paragraph = cell.paragraphs[0]
                _add_run(paragraph, value, settings, bold=(row_index < len(header_rows)))

                # Honor Markdown alignment when Mistune provides it.
                source_row = None
                if row_index < len(header_rows):
                    for node in block.get("children", []):
                        if node.get("type") == "table_head":
                            source_row = node.get("children", [])
                            break
                else:
                    body_rows = []
                    for node in block.get("children", []):
                        if node.get("type") == "table_body":
                            body_rows = [
                                row for row in node.get("children", [])
                                if row.get("type") == "table_row"
                            ]
                    body_index = row_index - len(header_rows)
                    if body_index < len(body_rows):
                        source_row = body_rows[body_index].get("children", [])

                if source_row and col_index < len(source_row):
                    align = source_row[col_index].get("attrs", {}).get("align")
                    paragraph.alignment = {
                        "left": WD_ALIGN_PARAGRAPH.LEFT,
                        "center": WD_ALIGN_PARAGRAPH.CENTER,
                        "right": WD_ALIGN_PARAGRAPH.RIGHT,
                    }.get(align, WD_ALIGN_PARAGRAPH.LEFT)

    def add_blocks(blocks, list_level=0):
        for block in blocks or []:
            kind = block.get("type")
            children = block.get("children", [])

            if kind == "blank_line":
                continue

            if kind == "heading":
                level = int(block.get("attrs", {}).get("level", 1))
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.keep_with_next = True
                paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
                paragraph.paragraph_format.space_after = Pt(5)
                for child in children:
                    if child.get("type") == "text":
                        _add_run(
                            paragraph, child.get("raw", ""), settings,
                            bold=True
                        )
                        for run in paragraph.runs[-1:]:
                            _apply_run_format(run, settings, heading=True, level=level)
                    else:
                        before = len(paragraph.runs)
                        inline_into(paragraph, [child])
                        for run in paragraph.runs[before:]:
                            _apply_run_format(run, settings, heading=True, level=level)
                continue

            if kind == "paragraph":
                paragraph = document.add_paragraph()
                inline_into(paragraph, children)
            elif kind == "block_quote":
                start = len(document.paragraphs)
                add_blocks(children, list_level)
                for paragraph in document.paragraphs[start:]:
                    paragraph.paragraph_format.left_indent = Cm(0.7)
            elif kind == "list":
                ordered = bool(block.get("attrs", {}).get("ordered", False))
                for item in children:
                    item_children = item.get("children", [])
                    style = "List Number" if ordered else "List Bullet"
                    paragraph = document.add_paragraph(style=style)
                    paragraph.paragraph_format.left_indent = Cm(0.6 + 0.45 * list_level)
                    first_block_written = False
                    for item_block in item_children:
                        item_kind = item_block.get("type")
                        if item_kind in ("block_text", "paragraph"):
                            if first_block_written:
                                paragraph = document.add_paragraph(style=style)
                                paragraph.paragraph_format.left_indent = Cm(0.6 + 0.45 * list_level)
                            inline_into(paragraph, item_block.get("children", []))
                            first_block_written = True
                        elif item_kind == "list":
                            add_blocks([item_block], list_level + 1)
                        elif item_kind == "block_code":
                            code = clean_xml_string(item_block.get("raw", "").rstrip("\n"))
                            code_paragraph = document.add_paragraph()
                            run = _add_run(code_paragraph, code, settings)
                            run.font.name = "Consolas"
            elif kind == "block_code":
                code = clean_xml_string(block.get("raw", "").rstrip("\n"))
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.5)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(8)
                run = _add_run(paragraph, code, settings)
                run.font.name = "Consolas"
                run.font.size = Pt(max(8, float(settings.get("body_size", 12)) - 1))
                run.font.color.rgb = _rgb("#374151")
            elif kind == "block_math":
                formula = clean_xml_string(block.get("raw", "").strip())
                if formula:
                    paragraph = document.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_before = Pt(4)
                    paragraph.paragraph_format.space_after = Pt(8)
                    run = _add_run(paragraph, formula, settings)
                    run.font.name = "Cambria Math"
                    run.italic = True
            elif kind == "thematic_break":
                paragraph = document.add_paragraph()
                _add_run(paragraph, "─" * 32, settings)
            elif kind == "table":
                add_table(block)
            elif kind == "block_html":
                raw = block.get("raw", "").strip()
                if raw:
                    paragraph = document.add_paragraph()
                    _add_inline(paragraph, raw, settings)
            else:
                text = plain_text(children) or block.get("raw", "")
                if text.strip():
                    paragraph = document.add_paragraph()
                    _add_inline(paragraph, text, settings)

    add_blocks(tokens)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
