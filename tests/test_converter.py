
from io import BytesIO

from docx import Document

from emydocx.converter import markdown_to_docx


def convert(markdown: str, settings: dict | None = None) -> Document:
    """Преобразовать Markdown и открыть полученный DOCX для проверки."""
    result = markdown_to_docx(markdown, settings or {})

    assert isinstance(result, bytes), "Конвертер должен возвращать bytes"
    assert len(result) > 0, "Полученный DOCX не должен быть пустым"

    return Document(BytesIO(result))


def all_paragraph_text(document: Document) -> str:
    """Собрать текст обычных абзацев документа."""
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_converts_simple_markdown_to_valid_docx():
    document = convert("# Привет\n\nЭто тестовый документ.")

    text = all_paragraph_text(document)

    assert "Привет" in text
    assert "Это тестовый документ." in text


def test_converts_heading():
    document = convert("# Главный заголовок")

    heading = next(
        paragraph
        for paragraph in document.paragraphs
        if "Главный заголовок" in paragraph.text
    )

    assert heading.text == "Главный заголовок"
    assert any(run.bold for run in heading.runs)


def test_converts_bold_and_italic_text():
    document = convert("Это **жирный** и *курсивный* текст.")

    runs = [
        run
        for paragraph in document.paragraphs
        for run in paragraph.runs
    ]

    bold_runs = [run for run in runs if run.bold]
    italic_runs = [run for run in runs if run.italic]

    assert any("жирный" in run.text for run in bold_runs)
    assert any("курсивный" in run.text for run in italic_runs)


def test_converts_bullet_list():
    document = convert("- Первый пункт\n- Второй пункт")

    list_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "List Bullet"
    ]

    assert len(list_paragraphs) == 2
    assert list_paragraphs[0].text == "Первый пункт"
    assert list_paragraphs[1].text == "Второй пункт"


def test_converts_numbered_list():
    document = convert("1. Первый пункт\n2. Второй пункт")

    list_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "List Number"
    ]

    assert len(list_paragraphs) == 2
    assert list_paragraphs[0].text == "Первый пункт"
    assert list_paragraphs[1].text == "Второй пункт"


def test_converts_table():
    markdown = (
        "| Имя | Возраст |\n"
        "| --- | ---: |\n"
        "| Анна | 25 |\n"
        "| Иван | 30 |"
    )

    document = convert(markdown)

    assert len(document.tables) == 1

    table = document.tables[0]

    assert len(table.rows) == 3
    assert len(table.columns) == 2
    assert table.cell(0, 0).text == "Имя"
    assert table.cell(0, 1).text == "Возраст"
    assert table.cell(1, 0).text == "Анна"
    assert table.cell(1, 1).text == "25"
    assert table.cell(2, 0).text == "Иван"
    assert table.cell(2, 1).text == "30"


def test_converts_code_block():
    document = convert(
        "```python\n"
        "print('Hello, world!')\n"
        "```"
    )

    assert "print('Hello, world!')" in all_paragraph_text(document)


def test_converts_link_to_label_and_url():
    document = convert("[Открыть сайт](https://example.com)")

    text = all_paragraph_text(document)

    assert "Открыть сайт" in text
    assert "https://example.com" in text


def test_removes_illegal_xml_control_characters():
    document = convert("До\x01после")

    text = all_paragraph_text(document)

    assert "До" in text
    assert "после" in text
    assert "\x01" not in text


def test_converts_russian_and_unicode_text():
    document = convert(
        "# Проверка Unicode\n\n"
        "Русский текст, Ўзбекистон, café, 中文."
    )

    text = all_paragraph_text(document)

    assert "Ўзбекистон" in text
    assert "café" in text
    assert "中文" in text



def test_empty_markdown_produces_valid_docx():
    document = convert("")

    assert document.core_properties is not None
    assert len(document.tables) == 0


def test_heading_uses_semantic_word_heading_style():
    document = convert("# Главный заголовок")

    heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "Главный заголовок"
    )

    assert heading.style.name == "Heading 1"


def test_nested_inline_formatting_preserves_bold_and_italic():
    document = convert("Это **жирный текст с *курсивом* внутри**.")

    target_runs = [
        run
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if "курсивом" in run.text
    ]

    assert target_runs, "Текст внутри форматирования не найден"
    assert any(run.bold and run.italic for run in target_runs)


def test_table_header_is_bold_but_body_is_not():
    from io import BytesIO
    from docx import Document

    markdown = (
        "| Название | Значение |\n"
        "| --- | --- |\n"
        "| Размер | 42 |"
    )

    result = markdown_to_docx(markdown, {})
    document = Document(BytesIO(result))
    table = document.tables[0]

    for cell in table.rows[0].cells:
        text_runs = [run for run in cell.paragraphs[0].runs if run.text]
        assert text_runs
        assert all(run.bold is True for run in text_runs)

    for cell in table.rows[1].cells:
        text_runs = [run for run in cell.paragraphs[0].runs if run.text]
        assert text_runs
        assert all(run.bold is False for run in text_runs)

def test_inline_code_uses_monospace_font():
    document = convert("Выполни `print('hello')`.")

    code_runs = [
        run
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if "print('hello')" in run.text
    ]

    assert code_runs, "Фрагмент кода не найден"
    assert all(run.font.name == "Consolas" for run in code_runs)


def test_nested_list_has_deeper_indentation():
    document = convert(
        "- Родительский пункт\n"
        "  - Вложенный пункт"
    )

    parent = next(
        paragraph
        for paragraph in document.paragraphs
        if "Родительский пункт" in paragraph.text
    )
    nested = next(
        paragraph
        for paragraph in document.paragraphs
        if "Вложенный пункт" in paragraph.text
    )

    assert nested.paragraph_format.left_indent is not None
    assert parent.paragraph_format.left_indent is not None
    assert (
        nested.paragraph_format.left_indent
        > parent.paragraph_format.left_indent
    )
