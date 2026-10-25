"""Блочная разметка: заголовки, списки, цитаты, код, таблицы."""
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm


# ---------- Заголовки ----------

@pytest.mark.parametrize("level", [1, 2, 3, 4, 5, 6])
def test_heading_style(level, settings, render):
    doc = render(f"{'#' * level} Title", settings)
    para = doc.paragraphs[0]
    assert para.style.name == f"Heading {level}"
    assert para.paragraph_format.keep_with_next is True


def test_h1_kept_size(settings, render):
    doc = render("# H", settings)
    run = doc.paragraphs[0].runs[0]
    assert run.font.size.pt == pytest.approx(settings["h1_size"])
    assert run.bold is True


def test_h4_falls_back_to_h3_size(settings, render):
    """Уровни 4+ используют h3_size — документируем поведение."""
    doc = render("#### H4", settings)
    run = doc.paragraphs[0].runs[0]
    assert run.font.size.pt == pytest.approx(settings["h3_size"])


# ---------- Параграфы ----------

def test_paragraph_creates_run(settings, render):
    doc = render("Hello world", settings)
    para = doc.paragraphs[0]
    assert para.text == "Hello world"
    assert para.runs[0].font.name == settings["font_name"]


# ---------- Списки ----------

def test_unordered_list(settings, render):
    doc = render("- a\n- b\n- c\n", settings)
    list_paras = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert len(list_paras) == 3
    assert [p.text for p in list_paras] == ["a", "b", "c"]


def test_ordered_list(settings, render):
    doc = render("1. a\n2. b\n", settings)
    list_paras = [p for p in doc.paragraphs if p.style.name == "List Number"]
    assert len(list_paras) == 2


def test_nested_list_indentation(settings, render):
    md = "- a\n  - b\n    - c\n"
    doc = render(md, settings)
    list_paras = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    indents = [p.paragraph_format.left_indent.cm for p in list_paras]
    assert len(indents) == 3
    assert indents[0] < indents[1] < indents[2]
    assert indents[0] == pytest.approx(0.6, abs=0.01)
    assert indents[1] == pytest.approx(1.05, abs=0.01)
    assert indents[2] == pytest.approx(1.5 , abs=0.01)


def test_mixed_list_styles(settings, render):
    md = "- a\n  1. b\n  2. c\n"
    doc = render(md, settings)
    bullets = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    numbers = [p for p in doc.paragraphs if p.style.name == "List Number"]
    assert len(bullets) == 1
    assert len(numbers) == 2


def test_list_item_with_code_block(settings, render):
    md = "- item\n\n  ```\n  code\n  ```\n"
    doc = render(md, settings)
    assert any("code" in p.text for p in doc.paragraphs)


# ---------- Цитаты ----------

def test_blockquote_indent(settings, render):
    doc = render("> quoted text", settings)
    para = doc.paragraphs[0]
    assert para.text == "quoted text"
    assert para.paragraph_format.left_indent.cm == pytest.approx(0.7, abs=0.01)


def test_blockquote_overrides_list_indent(settings, render):
    """Документируем: цитата перетирает indentation вложенного списка."""
    doc = render("> - a\n> - b\n", settings)
    list_paras = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert len(list_paras) == 2
    for p in list_paras:
        assert p.paragraph_format.left_indent.cm == pytest.approx(0.7, abs=0.01)


# ---------- Код ----------

def test_fenced_code_block(settings, render):
    md = "```python\nprint('hi')\n```\n"
    doc = render(md, settings)
    para = doc.paragraphs[0]
    assert "print('hi')" in para.text
    assert para.runs[0].font.name == "Consolas"
    assert para.paragraph_format.left_indent.cm == pytest.approx(0.5, abs=0.01)


def test_code_block_font_size_one_less_than_body(settings, render):
    doc = render("```\nx\n```", settings)
    run = doc.paragraphs[0].runs[0]
    assert run.font.size.pt == pytest.approx(settings["body_size"] - 1)


def test_code_block_strips_trailing_newlines(settings, render):
    doc = render("```\ncode\n\n\n```\n", settings)
    assert not doc.paragraphs[0].text.endswith("\n")


# ---------- Разделитель ----------

def test_thematic_break(settings, render):
    doc = render("---\n", settings)
    # Есть либо пустой параграф, либо параграф с символами ─
    assert any("─" in p.text for p in doc.paragraphs)


# ---------- Таблицы ----------

def test_table_basic_structure(settings, render):
    md = "| a | b | c |\n|---|---|---|\n| 1 | 2 | 3 |\n"
    doc = render(md, settings)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 2
    assert len(table.columns) == 3


def test_table_header_bold(settings, render):
    md = "| a | b |\n|---|---|\n| 1 | 2 |\n"
    doc = render(md, settings)
    table = doc.tables[0]
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            if run.text:
                assert run.bold is True


def test_table_body_not_bold(settings, render):
    md = "| a | b |\n|---|---|\n| 1 | 2 |\n"
    doc = render(md, settings)
    table = doc.tables[0]
    for cell in table.rows[1].cells:
        for run in cell.paragraphs[0].runs:
            if run.text:
                assert not run.bold


def test_table_alignment(settings, render):
    md = "| L | C | R |\n|:--|:-:|--:|\n| 1 | 2 | 3 |\n"
    doc = render(md, settings)
    body = doc.tables[0].rows[1]
    assert body.cells[0].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert body.cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert body.cells[2].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_uneven_table_row_via_markdown_does_not_crash(settings, render):
    """Mistune отвергает таблицу с разным числом ячеек — она становится
    текстом, но конвертер не падает. Поведение документируем."""
    doc = render("| a | b | c |\n|---|---|---|\n| 1 |\n", settings)
    # Либо нет таблицы вообще (mistune уронил её в paragraph),
    # либо таблица с padding — оба варианта валидны.
    # Главное — не упало и дало какой-то текст.
    all_text = " ".join(p.text for p in doc.paragraphs)
    assert "a" in all_text or len(doc.tables) >= 1


def test_renderer_pads_uneven_rows(settings):
    """_add_table сам добивает короткие строки — проверяем в изоляции
    от парсера (это и была цель выделения BlockRenderer)."""
    from docx import Document
    from emydocx.converter.renderer import BlockRenderer

    tokens = [{
        "type": "table",
        "children": [
            {"type": "table_head", "children": [
                {"type": "table_cell", "children": [{"type": "text", "raw": "a"}]},
                {"type": "table_cell", "children": [{"type": "text", "raw": "b"}]},
                {"type": "table_cell", "children": [{"type": "text", "raw": "c"}]},
            ]},
            {"type": "table_body", "children": [
                {"type": "table_row", "children": [
                    {"type": "table_cell", "children": [{"type": "text", "raw": "1"}]},
                ]},
            ]},
        ],
    }]

    doc = Document()
    BlockRenderer(doc, settings).render(tokens)

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 2
    assert len(table.rows[1].cells) == 3
    assert table.rows[1].cells[2].text == ""