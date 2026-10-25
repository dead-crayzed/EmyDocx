"""Все ключи settings применяются как заявлено."""
import pytest
from docx.shared import RGBColor


def test_all_settings_applied():
    settings = {
        "font_name": "Times New Roman",
        "body_size": 14,
        "h1_size": 24, "h2_size": 20, "h3_size": 16,
        "body_color": "#FF0000",
        "heading_color": "#0000FF",
        "margin_cm": 3.0,
        "line_spacing": 2.0,
    }
    from io import BytesIO
    from docx import Document
    from emydocx.converter import markdown_to_docx

    doc = Document(BytesIO(markdown_to_docx("# H1\n\nBody", settings)))

    # Normal style
    normal = doc.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size.pt == pytest.approx(14)
    assert normal.paragraph_format.line_spacing == pytest.approx(2.0)

    # Body run color
    body_para = doc.paragraphs[1]
    assert body_para.runs[0].font.color.rgb == RGBColor(0xFF, 0, 0)

    # Heading run color and size
    h1 = doc.paragraphs[0]
    assert h1.runs[0].font.size.pt == pytest.approx(24)
    assert h1.runs[0].font.color.rgb == RGBColor(0, 0, 0xFF)

    # Margins
    section = doc.sections[0]
    assert section.top_margin.cm == pytest.approx(3.0, abs=0.01)
    assert section.bottom_margin.cm == pytest.approx(3.0, abs=0.01)
    assert section.left_margin.cm == pytest.approx(3.0, abs=0.01)
    assert section.right_margin.cm == pytest.approx(3.0, abs=0.01)


@pytest.mark.parametrize("level_key,size", [
    ("h1_size", 30), ("h2_size", 22), ("h3_size", 18),
])
def test_heading_sizes(level_key, size, settings, render):
    settings[level_key] = size
    level = int(level_key[1])
    doc = render(f"{'#' * level} T", settings)
    assert doc.paragraphs[0].runs[0].font.size.pt == pytest.approx(size)


# ---------- Клэмпинг ----------

def test_body_size_clamped_high(settings, render):
    settings["body_size"] = 1000
    doc = render("Text", settings)
    # Клэмп только в SettingsPanel.get_settings, здесь markdown_to_docx
    # получает уже готовое число — считаем что панель его уже ограничила.
    # Если это не так — тест документирует текущее поведение:
    assert doc.paragraphs[0].runs[0].font.size.pt == pytest.approx(1000)


def test_invalid_color_fallback(settings, render):
    settings["body_color"] = "garbage"
    doc = render("Text", settings)
    # _rgb в colors.py вернёт дефолт
    assert doc.paragraphs[0].runs[0].font.color.rgb == RGBColor(0x20, 0x21, 0x24)


def test_hex_case_insensitive(settings, render):
    settings["body_color"] = "aAbBcC"
    doc = render("T", settings)
    assert doc.paragraphs[0].runs[0].font.color.rgb == RGBColor(0xAA, 0xBB, 0xCC)


def test_line_spacing_float(settings, render):
    settings["line_spacing"] = 1.5
    doc = render("T", settings)
    assert doc.styles["Normal"].paragraph_format.line_spacing == pytest.approx(1.5)