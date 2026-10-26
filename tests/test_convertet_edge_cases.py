"""Пустой ввод, странные символы, устойчивость."""
import pytest
from docx import Document
from io import BytesIO

from emydocx.converter import markdown_to_docx
from tests.conftest import list_properties


def test_empty_string(settings):
    out = markdown_to_docx("", settings)
    assert out[:2] == b"PK"
    Document(BytesIO(out))  # не падает


def test_none_input(settings):
    out = markdown_to_docx(None, settings)
    assert out[:2] == b"PK"


def test_whitespace_only(settings):
    out = markdown_to_docx("   \n\n\t\n", settings)
    assert out[:2] == b"PK"


def test_invalid_xml_chars_stripped(settings):
    out = markdown_to_docx("Hello\x00World\x01", settings)
    doc = Document(BytesIO(out))
    text = "".join(p.text for p in doc.paragraphs)
    assert "\x00" not in text
    assert "\x01" not in text
    assert "Hello" in text and "World" in text


def test_emoji_survive(settings):
    out = markdown_to_docx("Hello 🎉 and 🚀", settings)
    doc = Document(BytesIO(out))
    assert "🎉" in doc.paragraphs[0].text


def test_cjk_survives(settings):
    out = markdown_to_docx("Привет 世界", settings)
    doc = Document(BytesIO(out))
    assert "世界" in doc.paragraphs[0].text


def test_very_long_input(settings):
    md = "word " * 5000
    out = markdown_to_docx(md, settings)
    doc = Document(BytesIO(out))
    assert doc.paragraphs[0].text.count("word") == 5000


def test_deeply_nested_list(settings):
    md = "- a\n  - b\n    - c\n      - d\n        - e\n"
    out = markdown_to_docx(md, settings)
    doc = Document(BytesIO(out))
    list_paras = [p for p in doc.paragraphs if list_properties(doc, p)]
    assert len(list_paras) == 5
    properties = [list_properties(doc, p) for p in list_paras]
    assert [prop[1] for prop in properties] == list(range(5))
    assert [prop[2] for prop in properties] == ["bullet"] * 5
    assert [prop[3] for prop in properties] == [720, 1440, 2160, 2880, 3600]


def test_raw_html_block(settings):
    """HTML блок идёт в fallback regex-рендер."""
    out = markdown_to_docx("<div>raw</div>\n", settings)
    doc = Document(BytesIO(out))
    text = "".join(p.text for p in doc.paragraphs)
    assert "raw" in text or "div" in text


@pytest.mark.parametrize("md", [
    "# H\n\n**b** and *i* and `c`\n\n- x\n- y\n\n> q\n\n```\ncode\n```\n",
    "| a | b |\n|---|---|\n| 1 | 2 |\n\n---\n\nText",
    "1. one\n2. two\n\n![alt](img.png)\n\n[link](url)",
])
def test_kitchen_sink_does_not_crash(md, settings):
    out = markdown_to_docx(md, settings)
    Document(BytesIO(out))


def test_output_is_consistent_zip(settings):
    """Повторный вызов даёт валидный DOCX (ZIP magic)."""
    out1 = markdown_to_docx("# A", settings)
    out2 = markdown_to_docx("# A", settings)
    assert out1[:2] == b"PK"
    assert out2[:2] == b"PK"
